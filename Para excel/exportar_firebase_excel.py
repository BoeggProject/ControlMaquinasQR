import urllib.request
import requests
import json
import os
import shutil
from datetime import datetime, timedelta
from openpyxl import load_workbook

# =========================================================
# CONFIGURACIÓN
# =========================================================

BASE_URL = "https://piloto-614ca-default-rtdb.firebaseio.com/maquinas"

ARCHIVO_PLANTILLA = "Control_Maquinas_Semanal.xlsx"
ARCHIVO_PRUEBA_ANTERIOR = "Control_Maquinas_Semanal_actualizado.xlsx"
ARCHIVO_HISTORICO = "Control_Maquinas_Historico.xlsx"

# =========================================================
# DÍAS
# =========================================================

dias = {
    0: "LUNES",
    1: "MARTES",
    2: "MIERCOLES",
    3: "JUEVES",
    4: "VIERNES",
    5: "SABADO",
    6: "DOMINGO"
}

# =========================================================
# COLUMNAS DEL EXCEL
# =========================================================

columnas_dias = {
    0: 2,   # Lunes
    1: 5,   # Martes
    2: 8,   # Miércoles
    3: 11,  # Jueves
    4: 14,  # Viernes
    5: 17,  # Sábado
    6: 20   # Domingo
}

# =========================================================
# CONVERTIR FECHA Y HORA
# =========================================================

def convertir_fecha_hora(registro):
    fecha_texto = registro.get("fecha", "")
    hora_texto = registro.get("hora", "")
    if not fecha_texto:
        return None
    fecha = None
    formatos = ["%d/%m/%Y", "%d/%m/%y"]
    for formato in formatos:
        try:
            fecha = datetime.strptime(fecha_texto, formato)
            break
        except ValueError:
            pass
    if fecha is None:
        try:
            fecha = datetime.fromisoformat(fecha_texto.replace("Z", "+00:00"))
            fecha = fecha.replace(tzinfo=None)
        except:
            return None
    if not hora_texto:
        return fecha
    hora_limpia = hora_texto.lower().replace("a.m.", "AM").replace("p.m.", "PM").strip()
    try:
        hora = datetime.strptime(hora_limpia, "%I:%M:%S %p")
        return fecha.replace(hour=hora.hour, minute=hora.minute, second=hora.second)
    except:
        pass
    try:
        hora = datetime.strptime(hora_limpia, "%H:%M:%S")
        return fecha.replace(hour=hora.hour, minute=hora.minute, second=hora.second)
    except:
        pass
    return fecha

# =========================================================
# NOMBRE DE UNA SEMANA
# =========================================================

def nombre_semana(inicio, fin):
    meses = ["Ene","Feb","Mar","Abr","May","Jun","Jul","Ago","Sep","Oct","Nov","Dic"]
    mes_inicio = meses[inicio.month - 1]
    mes_fin = meses[fin.month - 1]
    if inicio.month == fin.month:
        return f"{inicio.day:02d}-{fin.day:02d} {mes_inicio}"
    return f"{inicio.day:02d} {mes_inicio}-{fin.day:02d} {mes_fin}"

# =========================================================
# OBTENER SEMANA ACTUAL
# =========================================================

hoy = datetime.now()
inicio_semana = hoy - timedelta(days=hoy.weekday())
inicio_semana = inicio_semana.replace(hour=0, minute=0, second=0, microsecond=0)
fin_semana = inicio_semana + timedelta(days=6, hours=23, minutes=59, seconds=59)
nombre_hoja_actual = nombre_semana(inicio_semana, fin_semana)

inicio_semana_anterior = inicio_semana - timedelta(days=7)
fin_semana_anterior = inicio_semana - timedelta(seconds=1)
nombre_hoja_anterior = nombre_semana(inicio_semana_anterior, fin_semana_anterior)

print()
print("==========================================")
print("   EXPORTADOR FIREBASE -> EXCEL")
print("==========================================")
print()
print("Semana actual:", inicio_semana.strftime("%d/%m/%Y"), "-", fin_semana.strftime("%d/%m/%Y"))
print("Hoja actual:", nombre_hoja_actual)

# =========================================================
# CREAR HISTÓRICO SI NO EXISTE
# =========================================================

if not os.path.exists(ARCHIVO_HISTORICO):
    print("\nCreando archivo histórico...")
    shutil.copy2(ARCHIVO_PLANTILLA, ARCHIVO_HISTORICO)

try:
    wb = load_workbook(ARCHIVO_HISTORICO)
except Exception as error:
    print("\nERROR al abrir el histórico:")
    print(error)
    input("\nPresiona Enter para cerrar...")
    exit()

# =========================================================
# CREAR / OBTENER HOJA ACTUAL
# =========================================================

if nombre_hoja_actual in wb.sheetnames:
    print("\nLa hoja actual ya existe. Se actualizará.")
    ws = wb[nombre_hoja_actual]
else:
    print("\nCreando hoja:", nombre_hoja_actual)
    hoja_base = wb[wb.sheetnames[0]]
    ws = wb.copy_worksheet(hoja_base)
    ws.title = nombre_hoja_actual

# =========================================================
# BUSCAR MÁQUINAS EN EXCEL
# =========================================================

filas_maquinas = {}
for fila in range(3, ws.max_row + 1):
    valor = ws.cell(row=fila, column=1).value
    if valor:
        id_maquina = str(valor).strip()
        filas_maquinas[id_maquina] = fila

print("\nMáquinas encontradas en Excel:", len(filas_maquinas))

# =========================================================
# CONECTAR FIREBASE (nuevo método)
# =========================================================

print("\nConectando con Firebase...")

datos_firebase = {}
for id_maquina in filas_maquinas.keys():
    url = f"{BASE_URL}/{id_maquina}/registros.json"
    try:
        resp = requests.get(url, timeout=20)
        if resp.status_code == 200:
            datos_firebase[id_maquina] = {"registros": resp.json()}
        else:
            print(f"Error leyendo {id_maquina}: {resp.status_code}")
    except Exception as error:
        print(f"Error conectando {id_maquina}: {error}")

print("Firebase conectado correctamente.")

# =========================================================
# RECORRER FIREBASE Y ESCRIBIR EN EXCEL
# =========================================================

total_actualizados = 0

for id_maquina, informacion in datos_firebase.items():
    print("\n------------------------------------------")
    print("Máquina:", id_maquina)
    print("------------------------------------------")

    if id_maquina not in filas_maquinas:
        print("  Esta máquina no está en el Excel. Se ignora.")
        continue

    registros = informacion.get("registros", {})
    if not registros:
        print("  No tiene registros.")
        continue

    registros_por_dia = {}
    for registro_id, registro in registros.items():
        if not isinstance(registro, dict):
            continue
        fecha_hora = convertir_fecha_hora(registro)
        if fecha_hora is None:
            continue
        if fecha_hora < inicio_semana or fecha_hora > fin_semana:
            continue
        dia_semana = fecha_hora.weekday()
        anterior = registros_por_dia.get(dia_semana)
        if anterior is None or fecha_hora > anterior[0]:
            registros_por_dia[dia_semana] = (fecha_hora, registro)

    if not registros_por_dia:
        print("  No tiene registros de esta semana.")
        continue

    fila_excel = filas_maquinas[id_maquina]
    for dia_semana, datos in registros_por_dia.items():
        fecha_hora, registro = datos
        columna = columnas_dias[dia_semana]
        nombre_dia = dias[dia_semana]
        temperatura = registro.get("temperatura", "")
        humedad = registro.get("humedad", "")
        porcentajeL = registro.get("porcentajeL", registro.get("luminosidad", ""))

        ws.cell(row=fila_excel, column=columna).value = temperatura
        ws.cell(row=fila_excel, column=columna + 1).value = humedad
        ws.cell(row=fila_excel, column=columna + 2).value = porcentajeL

        print(f"\n  Día: {nombre_dia}")
        print("  Último registro:", fecha_hora.strftime("%d/%m/%Y %H:%M:%S"))
        print(f"  → {temperatura} °C | {humedad} %H | {porcentajeL} %L")

        total_actualizados += 1

# =========================================================
# GUARDAR
# =========================================================

try:
    wb.save(ARCHIVO_HISTORICO)
except PermissionError:
    print("\nERROR: el Excel está abierto.")
    print